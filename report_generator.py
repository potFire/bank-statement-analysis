"""生成可追溯审核报告（DOCX/Markdown）与 JSON 明细。"""
from pathlib import Path
import json
from datetime import datetime,date
import pandas as pd
from docx import Document

def safe(value):
 if isinstance(value,(datetime,date,pd.Timestamp)):return str(value)
 if hasattr(value,"item"):return value.item()
 raise TypeError(str(type(value)))
def scrub(value):
 if isinstance(value,dict):return {k:scrub(v) for k,v in value.items()}
 if isinstance(value,list):return [scrub(v) for v in value]
 if isinstance(value,(datetime,date,pd.Timestamp)):return str(value)
 if hasattr(value,"item"):return value.item()
 return value

class ReportGenerator:
 def generate_report(self,metrics,company_name,output_dir,quality=None,findings=None):
  out=Path(output_dir);out.mkdir(parents=True,exist_ok=True)
  quality=quality or {};findings=findings or []
  metrics={k:v for k,v in metrics.items() if k!="classified_data"}
  payload=scrub({"company":company_name,"generated_at":datetime.now().isoformat(timespec="seconds"),"quality":quality,"metrics":metrics,"findings":findings,"disclaimer":"规则仅用于发现疑点，不构成银行流水真实性鉴定或授信审批结论；未做外部交叉验证。"})
  (out/"analysis.json").write_text(json.dumps(payload,ensure_ascii=False,indent=2,default=safe),encoding="utf-8")
  doc=Document();doc.add_heading("银行信贷流水审核分析报告",0)
  doc.add_paragraph("企业："+company_name)
  doc.add_paragraph(payload["disclaimer"])
  sections=[("数据质量",quality),("现金流",metrics.get("cash_flow",{})),("稳定性及季节性",metrics.get("stability",{}))]
  md=["# 银行信贷流水审核分析报告","","企业："+company_name,"",payload["disclaimer"],""]
  for name,content in sections:
   doc.add_heading(name,level=1);md.extend(["## "+name,""])
   for key,value in content.items():
    if isinstance(value,list):continue
    line=f"{key}: {value}"
    doc.add_paragraph(line);md.append("- "+line)
   md.append("")
  doc.add_heading("风险线索（非事实认定）",level=1)
  md.extend(["## 风险线索（非事实认定）",""])
  for f in findings[:300]:
   s=f"{f['level']} / {f['code']}: {f['reason']}，源行号 {f['source_rows']}，证据 {f['evidence']}"
   doc.add_paragraph(s,style="List Bullet");md.append("- "+s)
  if len(findings)>300:doc.add_paragraph("更多线索请查看 analysis.json")
  doc.add_heading("交易对手 TOP20",level=1)
  md.extend(["","## 交易对手 TOP20",""])
  for x in metrics.get("counterparties",[]):
   s=f"{x['name']} ({x['account']}): 流入 {x['inflow']}，流出 {x['outflow']}，笔数 {x['count']}"
   doc.add_paragraph(s);md.append("- "+s)
  doc.save(out/"analysis.docx")
  (out/"analysis.md").write_text("\n".join(md)+"\n",encoding="utf-8")
  return str(out/"analysis.docx")
