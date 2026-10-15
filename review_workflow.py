"""人工复核交付：生成 Excel 复核清单，全部确认后重新导入再分析。

注意：表格“已确认”仅表示操作人已复核，没有身份签名、机构授权。
"""
from datetime import datetime, timezone
from pathlib import Path
import hashlib
import json
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from statement_schema import STANDARD_COLUMNS, validate_standard_frame

REVIEW_STATUS = "审核状态"
REVIEW_NOTE = "审核备注"
APPROVED_VALUES = {"已确认", "通过", "approved"}


def _lineage_digest(data):
    """来源页、工作表、行号不可在复核过程中悄然删改。"""
    values = [
        "|".join((
            str(r["source_file"]), str(r["source_sheet"]),
            str(int(r["source_page"])), str(int(r["source_row"]))
        ))
        for _, r in data.iterrows()
    ]
    return hashlib.sha256("\n".join(sorted(values)).encode("utf-8")).hexdigest()


def export_review_workbook(data, origins, output_dir):
    validate_standard_frame(data)
    folder = Path(output_dir)
    folder.mkdir(parents=True, exist_ok=True)
    outfile = folder / "recognition_review.xlsx"
    manifest = folder / "recognition_manifest.json"
    book = Workbook()
    sheet = book.active
    sheet.title = "标准流水复核"
    headings = list(STANDARD_COLUMNS) + [REVIEW_STATUS, REVIEW_NOTE]
    sheet.append(headings)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{sheet.cell(1, len(headings)).column_letter}{len(data) + 1}"
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(fill_type="solid", fgColor="DDEBF7")
    for _, row in data.iterrows():
        values = []
        for col in STANDARD_COLUMNS:
            value = row[col]
            if pd.isna(value):
                value = ""
            elif isinstance(value, pd.Timestamp):
                value = value.isoformat()
            values.append(value)
        values.extend(["待复核" if bool(row["review_required"]) else "已确认", ""])
        sheet.append(values)
        for col_index, column in enumerate(STANDARD_COLUMNS, start=1):
            # 所有文本严格写成 Excel 字符串，避免原始流水摘要被当公式执行。
            cell = sheet.cell(sheet.max_row, col_index)
            if column not in ("贷方发生额", "借方发生额", "余额",
                              "source_row", "source_page", "recognition_score",
                              "review_required") and isinstance(cell.value, str):
                cell.data_type = "s"
                cell.number_format = "@"
        if values[-2] == "待复核":
            sheet.cell(sheet.max_row, len(headings) - 1).fill = PatternFill(
                fill_type="solid", fgColor="FFF2CC"
            )
    book.save(outfile)
    payload = {
        "schema_version": "1.1",
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "rows": int(len(data)),
        "require_review": int(data["review_required"].fillna(False).astype(bool).sum()),
        "sources": origins,
        "workbook_sha256": hashlib.sha256(outfile.read_bytes()).hexdigest(),
        "origin_lineage_sha256": _lineage_digest(data),
        "note": "需逐笔核对原件。将全部审核状态改为 已确认 后，才可使用 --reviewed-xlsx 导入；本文件不具备数字签名功能。",
    }
    manifest.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return outfile


def read_approved_workbook(path):
    book = load_workbook(path, read_only=False, data_only=False)
    if "标准流水复核" not in book:
        raise ValueError("复核文件缺少“标准流水复核”工作表")
    sheet = book["标准流水复核"]
    rows = list(sheet.iter_rows(values_only=False))
    if not rows:
        raise ValueError("复核文件为空")
    headers = [cell.value for cell in rows[0]]
    expected = list(STANDARD_COLUMNS) + [REVIEW_STATUS, REVIEW_NOTE]
    if headers != expected:
        raise ValueError("复核文件的列结构被修改，请保持标准字段顺序")
    records = []
    for excel_row in rows[1:]:
        if all(c.value is None for c in excel_row):
            continue
        if any(cell.data_type == "f" for cell in excel_row):
            raise ValueError("复核表格不能包含 Excel 公式")
        values = [cell.value for cell in excel_row]
        status = str(values[-2] or "").strip().lower()
        if status not in APPROVED_VALUES:
            raise ValueError(f"复核表第 {excel_row[0].row} 行尚未标记 已确认")
        records.append(values[:len(STANDARD_COLUMNS)])
    if not records:
        raise ValueError("复核表没有交易记录")
    data = pd.DataFrame(records, columns=STANDARD_COLUMNS)
    manifest = Path(path).with_name("recognition_manifest.json")
    if not manifest.is_file():
        raise ValueError("缺少识别清单对应的 recognition_manifest.json，无法验证来源完整性")
    metadata = json.loads(manifest.read_text(encoding="utf-8"))
    if len(data) != metadata.get("rows"):
        raise ValueError("复核交易行数与原始识别清单不一致；不可直接删除未确认交易")
    if _lineage_digest(data) != metadata.get("origin_lineage_sha256"):
        raise ValueError("复核清单的原始文件、页码或行号被修改，拒绝进入审核")
    # 人工逐笔确认后解除待审核标记；保留来源方式与结构评分。
    data["review_required"] = False
    validate_standard_frame(data)
    return data
