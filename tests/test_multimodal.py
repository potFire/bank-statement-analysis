"""脱敏的合成 PDF/图片样本与模拟视觉模型测试，不调用外部 API。"""
from pathlib import Path
from io import BytesIO
import fitz
from PIL import Image
import pytest
from openpyxl import load_workbook

from main import analyze
from data_processor import DataProcessor
from review_workflow import read_approved_workbook
from vision_llm import LLMSettings, validate_extraction


COLS = ["交易日期", "账号", "贷方发生额", "借方发生额", "余额", "交易用途", "对方账号"]
ROWS = [
    ["2026-03-01", "00123", "90000", "0", "90000", "货款", "OPPOSITE"],
    ["2026-03-02", "00123", "0", "20000", "70000", "采购", "OPPOSITE"],
]


class FakeVisionLLM:
    def __init__(self):
        self.image_calls = 0
        self.text_calls = 0

    def extract_image(self, image_bytes, media_type="image/png"):
        self.image_calls += 1
        assert image_bytes[:8] == b"\x89PNG\r\n\x1a\n"
        return {"columns": COLS, "rows": ROWS, "bank_hint": "中国农业银行"}

    def extract_text(self, text):
        self.text_calls += 1
        assert len(text) >= 100
        return {"columns": COLS, "rows": ROWS, "bank_hint": "中国农业银行"}


def test_image_roundtrip_mandatory_review(tmp_path):
    path = tmp_path / "statement.png"
    Image.new("RGB", (640, 300), "white").save(path)
    model = FakeVisionLLM()
    target = tmp_path / "result"
    preview = Path(analyze(path, "测试企业", target, llm_client=model))
    assert preview.exists() and preview.name == "recognition_review.xlsx"
    assert model.image_calls == 1
    assert not (target / "analysis.docx").exists()

    book = load_workbook(preview)
    tab = book["标准流水复核"]
    assert tab.max_row == 3
    header = [c.value for c in tab[1]]
    assert "source_page" in header and "extraction_method" in header
    assert tab.cell(2, header.index("账号") + 1).value == "00123"
    assert tab.cell(2, header.index("source_page") + 1).value == 1
    for row in (2, 3):
        tab.cell(row, header.index("审核状态") + 1, "已确认")
    book.save(preview)
    reviewed = read_approved_workbook(preview)
    assert len(reviewed) == 2
    result = Path(analyze(None, "测试企业", target, reviewed_xlsx=preview))
    assert result.exists()
    text = (target / "analysis.md").read_text(encoding="utf-8")
    assert "user_attested_approved" in text
    assert "未经人工复核" not in text


def test_image_cannot_be_auto_analyzed_without_model(tmp_path):
    path = tmp_path / "scan.jpg"
    Image.new("RGB", (200, 100), "white").save(path)
    with pytest.raises(ValueError, match="本地视觉 LLM"):
        DataProcessor().process_all(path)


def test_unreviewed_draft_is_explicit(tmp_path):
    path = tmp_path / "scan.png"
    Image.new("RGB", (100, 100), "white").save(path)
    result = Path(analyze(
        path, "测试企业", tmp_path / "out",
        llm_client=FakeVisionLLM(), allow_unreviewed=True,
    ))
    assert result.exists()
    assert "未经人工复核" in (tmp_path / "out" / "analysis.md").read_text(encoding="utf-8")


def test_scanned_pdf_vision_flow(tmp_path):
    image = Image.new("RGB", (600, 400), "white")
    buf = BytesIO()
    image.save(buf, "PNG")
    document = fitz.open()
    page = document.new_page(width=600, height=400)
    page.insert_image(fitz.Rect(0, 0, 600, 400), stream=buf.getvalue())
    path = tmp_path / "scan.pdf"
    document.save(path)
    document.close()
    stub = FakeVisionLLM()
    data, report = DataProcessor(llm_client=stub).process_all(path)
    assert len(data) == 2
    assert set(data["source_page"]) == {1}
    assert set(data["extraction_method"]) == {"llm_pdf_vision"}
    assert stub.image_calls == 1
    assert report["review_required_records"] == 2


def test_text_based_pdf_also_uses_vision_llm(tmp_path):
    document = fitz.open()
    page = document.new_page()
    page.insert_textbox(
        fitz.Rect(30, 30, 500, 350),
        ("Bank account statement records transaction 2026-03-01 "
         "credit 90000 debit 0 balance 90000. ") * 4,
        fontsize=11,
    )
    path = tmp_path / "text.pdf"
    document.save(path)
    document.close()
    stub = FakeVisionLLM()
    data, report = DataProcessor(llm_client=stub).process_all(path)
    assert len(data) == 2
    assert set(data["extraction_method"]) == {"llm_pdf_vision"}
    assert stub.image_calls == 1
    assert stub.text_calls == 0


def test_text_layer_pdf_table_requires_vision_llm(tmp_path):
    # 即使 PDF 含文字表格和线框，也不能直接用 PyMuPDF 提取绕过大模型。
    cells = [
        ["date", "account", "credit", "debit", "balance", "purpose"],
        ["2026-03-01", "00123", "100", "0", "100", "payment"],
        ["2026-03-02", "00123", "0", "30", "70", "supply"],
    ]
    widths = [105, 75, 70, 70, 80, 100]
    xpos = [30]
    for width in widths:
        xpos.append(xpos[-1] + width)
    ypos = [60, 89, 118, 147]
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    for x in xpos:
        page.draw_line((x, ypos[0]), (x, ypos[-1]))
    for y in ypos:
        page.draw_line((xpos[0], y), (xpos[-1], y))
    for row_idx, row in enumerate(cells):
        for col_idx, value in enumerate(row):
            page.insert_text((xpos[col_idx] + 4, ypos[row_idx] + 18), value, fontsize=8)
    path = tmp_path / "typed.pdf"
    document.save(path)
    document.close()
    with pytest.raises(ValueError, match="必须通过视觉大模型"):
        DataProcessor().process_all(path)

    class TableVision(FakeVisionLLM):
        def extract_image(self, image_bytes, media_type="image/png"):
            super().extract_image(image_bytes, media_type)
            return {"columns": cells[0], "rows": cells[1:], "bank_hint": ""}

    model = TableVision()
    data, info = DataProcessor(llm_client=model).process_all(path)
    assert len(data) == 2
    assert set(data["extraction_method"]) == {"llm_pdf_vision"}
    assert data["贷方发生额"].tolist() == [100, 0]
    assert info["balance_error_count"] == 0
    assert model.image_calls == 1
    assert model.text_calls == 0


def test_bad_vision_output_and_remote_protection():
    with pytest.raises(ValueError, match="远端"):
        LLMSettings(url="https://provider.example/v1/chat/completions", model="vision")
    with pytest.raises(ValueError, match="HTTPS"):
        LLMSettings(url="http://provider.example/v1/chat/completions",
                    model="vision", allow_remote=True)
    local = LLMSettings(url="http://127.0.0.1:11434/v1/chat/completions", model="local")
    assert local.model == "local"
    with pytest.raises(ValueError, match="列数"):
        validate_extraction({"columns": ["日期", "金额"], "rows": [["2026-01-01"]]})


def test_pdf_page_limit(tmp_path):
    pdf = fitz.open()
    pdf.new_page()
    pdf.new_page()
    path = tmp_path / "twopages.pdf"
    pdf.save(path)
    pdf.close()
    with pytest.raises(ValueError, match="超过最大 1 页"):
        DataProcessor(max_pages=1, llm_client=FakeVisionLLM()).process_all(path)


def test_model_unreadable_amount_is_kept_for_manual_correction(tmp_path):
    class UnclearModel(FakeVisionLLM):
        def extract_image(self, image_bytes, media_type="image/png"):
            super().extract_image(image_bytes, media_type)
            wrong = [row.copy() for row in ROWS]
            wrong[0][2] = "无法辨认"
            return {"columns": COLS, "rows": wrong, "bank_hint": "中国农业银行"}

    image = tmp_path / "unclear.png"
    Image.new("RGB", (300, 180), "white").save(image)
    review = Path(analyze(image, "样例", tmp_path / "out", llm_client=UnclearModel()))
    book = load_workbook(review)
    tab = book["标准流水复核"]
    headers = [cell.value for cell in tab[1]]
    amount_col = headers.index("贷方发生额") + 1
    status_col = headers.index("审核状态") + 1
    assert tab.cell(2, amount_col).value == "无法辨认"
    tab.cell(2, status_col, "已确认")
    tab.cell(3, status_col, "已确认")
    book.save(review)
    with pytest.raises(ValueError, match="仍存在无法解析"):
        analyze(None, "样例", tmp_path / "out", reviewed_xlsx=review)
    book = load_workbook(review)
    tab = book["标准流水复核"]
    tab.cell(2, amount_col, "90000.00")
    book.save(review)
    output = Path(analyze(None, "样例", tmp_path / "out", reviewed_xlsx=review))
    assert output.exists()


def test_source_lineage_rejects_removed_transaction(tmp_path):
    photo = tmp_path / "scan.png"
    Image.new("RGB", (300, 160), "white").save(photo)
    review = Path(analyze(photo, "样例", tmp_path / "out", llm_client=FakeVisionLLM()))
    book = load_workbook(review)
    tab = book["标准流水复核"]
    col = [cell.value for cell in tab[1]].index("审核状态") + 1
    tab.cell(2, col, "已确认")
    tab.delete_rows(3)
    book.save(review)
    with pytest.raises(ValueError, match="行数"):
        read_approved_workbook(review)


def test_source_lineage_rejects_rewritten_page_or_row(tmp_path):
    photo = tmp_path / "scan.png"
    Image.new("RGB", (300, 160), "white").save(photo)
    review = Path(analyze(photo, "样例", tmp_path / "out", llm_client=FakeVisionLLM()))
    book = load_workbook(review)
    tab = book["标准流水复核"]
    header = [cell.value for cell in tab[1]]
    status_col = header.index("审核状态") + 1
    row_col = header.index("source_row") + 1
    for row in (2, 3):
        tab.cell(row, status_col, "已确认")
    tab.cell(2, row_col, 999)
    book.save(review)
    with pytest.raises(ValueError, match="原始文件、页码或行号"):
        read_approved_workbook(review)


def test_openai_compatible_llm_payload_without_network(monkeypatch):
    import json
    import vision_llm
    from vision_llm import VisionLLM
    response_content = json.dumps({
        "columns": COLS, "rows": ROWS, "bank_hint": "农行",
    }, ensure_ascii=False)
    observed = {}

    def fake_urlopen(req, timeout):
        observed["body"] = json.loads(req.data.decode("utf-8"))
        observed["authorization"] = req.get_header("Authorization")
        observed["timeout"] = timeout
        return BytesIO(json.dumps({
            "choices": [{"message": {"content": response_content}}]
        }).encode("utf-8"))

    monkeypatch.setattr(vision_llm.request, "urlopen", fake_urlopen)
    monkeypatch.setenv("BANK_LLM_API_KEY", "test-secret")
    client = VisionLLM(LLMSettings(
        url="http://127.0.0.1:11434/v1/chat/completions",
        model="test-vision-model",
    ))
    result = client.extract_image(b"synthetic-image-bytes")
    assert result["rows"] == ROWS
    assert observed["body"]["model"] == "test-vision-model"
    assert observed["body"]["messages"][1]["content"][1]["type"] == "image_url"
    assert observed["authorization"] == "Bearer test-secret"
    assert observed["timeout"] == 180


def test_pdf_every_page_including_blank_is_processed_by_vision(tmp_path):
    doc = fitz.open()
    doc.new_page().insert_text((40, 80), "Bank statement: page one")
    doc.new_page()  # 空白页也交给模型判别，不根据文字层决定路由
    doc.new_page().insert_text((40, 80), "Bank statement: page three")
    path = tmp_path / "multipage.pdf"
    doc.save(path)
    doc.close()

    class PageVision(FakeVisionLLM):
        def extract_image(self, image_bytes, media_type="image/png"):
            super().extract_image(image_bytes, media_type)
            if self.image_calls == 2:
                return {"columns": [], "rows": [], "notes": "空白页"}
            return {"columns": COLS, "rows": ROWS, "bank_hint": "中国农业银行"}

    client = PageVision()
    frame, quality = DataProcessor(llm_client=client).process_all(path)
    assert client.image_calls == 3
    assert client.text_calls == 0
    assert len(frame) == 4
    assert sorted(frame["source_page"].unique().tolist()) == [1, 3]
    assert len(quality["sources"]) == 3
    assert all(item["extraction_method"] == "llm_pdf_vision" for item in quality["sources"])


def test_pdf_page_failure_cannot_silently_skip(tmp_path):
    doc = fitz.open()
    doc.new_page().insert_text((40, 80), "Bank statement page one")
    doc.new_page().insert_text((40, 80), "Bank statement page two")
    path = tmp_path / "fail.pdf"
    doc.save(path)
    doc.close()

    class FailingVision(FakeVisionLLM):
        def extract_image(self, image_bytes, media_type="image/png"):
            if self.image_calls == 1:
                raise RuntimeError("synthetic model failure")
            return super().extract_image(image_bytes, media_type)

    with pytest.raises(ValueError, match="第 2 页"):
        DataProcessor(llm_client=FailingVision()).process_all(path)
